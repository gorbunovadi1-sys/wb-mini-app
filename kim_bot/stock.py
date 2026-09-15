"""Baseline stock upload/parsing and the daily balance calculation.

Balance for an article = baseline qty (as of the date she counted it)
  + returns registered since that date (any status — see project decision)
  - orders created since that date that weren't cancelled.

Orders count against stock from creation, not from confirmed-packed —
simpler and matches how sellers usually think of "reserved" stock; daily-
granularity reconciliation means the few hours between order and pack don't
matter here the way they do for the separate SLA alert."""
import datetime
import io

import openpyxl

from .db import SessionLocal
from .models import DailyStockSnapshot, FbsOrder, ReturnRecord, StockBaseline

TEMPLATE_HEADER = ["Артикул", "Остаток"]
DATE_LABEL = "Дата остатка (ГГГГ-ММ-ДД):"


def build_empty_template() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Остаток"
    ws["A1"] = DATE_LABEL
    ws["B1"] = datetime.date.today().isoformat()
    ws.append([])
    ws.append(TEMPLATE_HEADER)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 16
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def parse_template(file_bytes: bytes) -> tuple[str, dict]:
    """Returns (as_of_date iso string, {article: qty}). Raises ValueError on
    a missing/unparseable date or no data rows."""
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    ws = wb.active

    raw_date = ws["B1"].value
    if raw_date is None:
        raise ValueError(f"Не нашла дату в B1 — первая строка должна быть «{DATE_LABEL}» с датой рядом.")
    as_of_date = raw_date.isoformat() if hasattr(raw_date, "isoformat") else str(raw_date).strip()[:10]
    try:
        datetime.date.fromisoformat(as_of_date)
    except ValueError:
        raise ValueError(f"Не смогла разобрать дату «{raw_date}» в B1 — нужен формат ГГГГ-ММ-ДД.")

    qty_by_article = {}
    for row in ws.iter_rows(min_row=4, values_only=True):
        if not row or row[0] is None:
            continue
        article = str(row[0]).strip()
        try:
            qty = int(row[1])
        except (TypeError, ValueError):
            continue
        if article:
            qty_by_article[article] = qty

    if not qty_by_article:
        raise ValueError("Не нашла строк с артикулами и остатками — проверь файл (данные с 4-й строки).")

    return as_of_date, qty_by_article


def replace_baseline(as_of_date: str, qty_by_article: dict):
    """Wholesale replace — a new upload fully supersedes the old baseline,
    matching how she described re-sending a fresh count."""
    with SessionLocal() as db:
        db.query(StockBaseline).delete()
        for article, qty in qty_by_article.items():
            db.add(StockBaseline(article=article, qty=qty, as_of_date=as_of_date))
        db.commit()


def get_baseline_date():
    """Earliest as_of_date currently in the baseline, if any — used to know
    how far back order/return polling needs to look to cover it."""
    with SessionLocal() as db:
        dates = [b.as_of_date for b in db.query(StockBaseline.as_of_date).distinct()]
        return min(dates) if dates else None


def compute_balances() -> dict:
    """Returns {article: balance} for every article in the current baseline."""
    with SessionLocal() as db:
        baselines = db.query(StockBaseline).all()
        balances = {}
        for b in baselines:
            since = datetime.datetime.fromisoformat(b.as_of_date)
            returned = sum(
                r.qty for r in db.query(ReturnRecord).filter(
                    ReturnRecord.article == b.article, ReturnRecord.created_at >= since,
                )
            )
            shipped = sum(
                o.qty for o in db.query(FbsOrder).filter(
                    FbsOrder.article == b.article, FbsOrder.created_at >= since,
                    FbsOrder.cancelled_at.is_(None),
                )
            )
            balances[b.article] = b.qty + returned - shipped
        return balances


def build_summary_xlsx(balances: dict) -> bytes:
    """Артикул | Остаток, sorted ascending by balance — lowest stock (the
    stuff worth acting on) shows up first instead of being buried alphabetically."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Остаток"
    ws.append(["Артикул", "Остаток"])
    for article, qty in sorted(balances.items(), key=lambda kv: kv[1]):
        ws.append([article, qty])
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 14
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def save_daily_snapshot(balances: dict, snapshot_date: str = None):
    snapshot_date = snapshot_date or datetime.date.today().isoformat()
    with SessionLocal() as db:
        for article, balance in balances.items():
            existing = db.query(DailyStockSnapshot).filter_by(article=article, snapshot_date=snapshot_date).first()
            if existing:
                existing.balance = balance
                existing.computed_at = datetime.datetime.utcnow()
            else:
                db.add(DailyStockSnapshot(article=article, snapshot_date=snapshot_date, balance=balance))
        db.commit()
