"""Automated version of her manual "Сверка_отгрузок_WB_Ozon.xlsx" — compares
what our own API polling (orders.py) recorded against what fulfillment
entered by hand in their Google Sheet, and flags where they disagree, with
a plain-language comment per finding (which platform, what kind of gap).
Ozon FBO is shown journal-only (we don't poll FBO postings at all), same
limitation her manual version had."""
import collections
import datetime
import io
import logging

import openpyxl
from openpyxl.styles import Font, PatternFill

from . import fulfillment_sheet, stock
from .db import SessionLocal
from .models import FbsOrder

log = logging.getLogger("kim_bot.reconciliation")

RED = PatternFill("solid", fgColor="FFD9D9")
BOLD = Font(bold=True)
PLATFORM_LABELS = {"wb": "WB", "ozon": "Ozon"}


MSK_OFFSET = datetime.timedelta(hours=3)


def _api_shipments(days: int) -> list[dict]:
    """Every row here is FBS — we only ever poll FBS postings/orders.
    created_at is stored as naive UTC (see util.parse_dt); her journal and
    manual audits date things by Moscow day, so dates here must be shifted
    to MSK before bucketing — an order at 21:47 UTC is already the next day
    in Moscow, and comparing raw UTC dates against her journal silently
    shifted a chunk of every evening's orders onto the wrong day (verified
    live 2026-09-15 against her own 15.08 audit: 15 raw-UTC vs the correct 12
    once shifted to MSK, matching her manual count exactly)."""
    since = datetime.datetime.utcnow() - datetime.timedelta(days=days)
    with SessionLocal() as db:
        rows = db.query(FbsOrder).filter(
            FbsOrder.created_at >= since, FbsOrder.cancelled_at.is_(None), FbsOrder.article.isnot(None),
        ).all()
        return [
            {"date": (o.created_at + MSK_OFFSET).date(), "platform": o.marketplace, "scheme": "FBS", "article": o.article, "qty": o.qty}
            for o in rows
        ]


def _mismatch_comment(platform: str, api_qty: int, j_qty: int) -> str:
    label = PLATFORM_LABELS.get(platform, platform)
    if j_qty == 0:
        return f"{label}: в журнале ФФ нет записи вообще, хотя по API отгружено {api_qty} шт — скорее всего забыли занести"
    if api_qty == 0:
        return f"{label}: запись в журнале ФФ есть ({j_qty} шт), но в API за эту дату/артикул её нет — проверить дату или артикул в журнале"
    if api_qty > j_qty:
        return f"{label}: журнал ФФ занижает — по факту {api_qty} шт, записано только {j_qty}"
    return f"{label}: журнал ФФ завышает — по факту {api_qty} шт, записано {j_qty} (возможен задвоенный ввод)"


def _find_full_day_gaps(api_rows: list, journal_rows: list) -> list[str]:
    """Whole-day blackouts — a platform had real API activity on a date but
    zero journal entries for it at all — reported as one consecutive-range
    finding per platform, matching the kind of gap her manual sverka found
    for WB 22–31.08 (one big missing stretch, not 10 separate line items)."""
    api_dates_by_platform = collections.defaultdict(set)
    for r in api_rows:
        api_dates_by_platform[r["platform"]].add(r["date"])
    journal_dates_by_platform = collections.defaultdict(set)
    for r in journal_rows:
        journal_dates_by_platform[r["platform"]].add(r["date"])

    findings = []
    for platform, dates in api_dates_by_platform.items():
        missing = sorted(d for d in dates if d not in journal_dates_by_platform.get(platform, set()))
        if not missing:
            continue
        # collapse consecutive dates into ranges
        ranges = []
        start = prev = missing[0]
        for d in missing[1:]:
            if (d - prev).days == 1:
                prev = d
                continue
            ranges.append((start, prev))
            start = prev = d
        ranges.append((start, prev))
        label = PLATFORM_LABELS.get(platform, platform)
        for start, end in ranges:
            span = f"{start.isoformat()}" if start == end else f"{start.isoformat()}–{end.isoformat()}"
            findings.append(f"{label}: в журнале ФФ нет ни одной записи за {span}, хотя по API в эти дни были отгрузки")
    return findings


def build_report(days: int = 14) -> bytes:
    api_rows = _api_shipments(days)
    cutoff = datetime.date.today() - datetime.timedelta(days=days)
    journal_rows = [j for j in fulfillment_sheet.parse_journal_entries() if j["date"] >= cutoff]

    api_totals = collections.Counter()
    for r in api_rows:
        api_totals[(r["platform"], r["scheme"])] += r["qty"]
    journal_totals = collections.Counter()
    for r in journal_rows:
        journal_totals[(r["platform"], r["scheme"])] += r["qty"]
    combos = sorted(set(api_totals) | set(journal_totals) | {("ozon", "FBO")})

    api_by_key = collections.Counter()
    for r in api_rows:
        api_by_key[(r["article"], r["date"], r["platform"], r["scheme"])] += r["qty"]
    journal_by_key = collections.Counter()
    for r in journal_rows:
        journal_by_key[(r["article"], r["date"], r["platform"], r["scheme"])] += r["qty"]

    mismatches = []
    for key in set(api_by_key) | set(journal_by_key):
        article, date, platform, scheme = key
        if platform == "ozon" and scheme == "FBO":
            continue  # no API fact to compare against — see module docstring
        api_qty, j_qty = api_by_key.get(key, 0), journal_by_key.get(key, 0)
        if api_qty != j_qty:
            mismatches.append((article, date, platform, scheme, api_qty, j_qty))
    mismatches.sort(key=lambda m: m[1])

    # --- narrative findings for the top of "Сводка" ---
    findings = []
    for platform, scheme in combos:
        if platform == "ozon" and scheme == "FBO":
            continue
        api_qty, j_qty = api_totals.get((platform, scheme), 0), journal_totals.get((platform, scheme), 0)
        if api_qty != j_qty:
            findings.append(
                f"{PLATFORM_LABELS.get(platform, platform)} {scheme}: недоучтено в журнале {api_qty - j_qty} шт "
                f"(факт {api_qty}, журнал {j_qty})" if api_qty > j_qty else
                f"{PLATFORM_LABELS.get(platform, platform)} {scheme}: в журнале больше, чем по факту, на {j_qty - api_qty} шт "
                f"(факт {api_qty}, журнал {j_qty})"
            )
    findings.extend(_find_full_day_gaps(api_rows, journal_rows))
    if not findings:
        findings = ["Расхождений факт vs журнал ФФ не найдено — данные сходятся."]

    balances = stock.compute_balances()
    ff_stock = fulfillment_sheet.fetch_current_ff_stock()
    names_by_article = {
        (row.get("Артикул завода") or "").strip(): (row.get("Наименование") or "").strip()
        for row in fulfillment_sheet.fetch_article_rows()
    }

    wb = openpyxl.Workbook()

    # --- Сводка ---
    ws = wb.active
    ws.title = "Сводка"
    ws.append([f"Автосверка отгрузок WB + Ozon — последние {days} дн."])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Сформировано: {datetime.date.today().isoformat()}"])
    ws.append([])
    ws.append(["Площадка", "Схема", "Отгружено по факту (API)", "Отгружено по журналу ФФ", "Расхождение"])
    for c in ws[ws.max_row]:
        c.font = BOLD
    for platform, scheme in combos:
        no_api_fact = platform == "ozon" and scheme == "FBO"
        api_qty = api_totals.get((platform, scheme), 0)
        j_qty = journal_totals.get((platform, scheme), 0)
        diff = "—" if no_api_fact else api_qty - j_qty
        ws.append([PLATFORM_LABELS.get(platform, platform), scheme, "нет данных из API" if no_api_fact else api_qty, j_qty, diff])
        if isinstance(diff, int) and diff != 0:
            ws.cell(ws.max_row, 5).fill = RED
    ws.append([])
    ws.append(["Главные находки"])
    ws.cell(ws.max_row, 1).font = BOLD
    for f in findings:
        ws.append([f])
    for col, width in zip("ABCDE", (10, 10, 24, 24, 14)):
        ws.column_dimensions[col].width = width

    # --- Остатки — сходимость ---
    ws2 = wb.create_sheet("Остатки — сходимость")
    ws2.append(["Артикул", "Остаток (бот)", "Остаток по журналу ФФ", "Расхождение"])
    for c in ws2[1]:
        c.font = BOLD
    for article in sorted(set(balances) | set(ff_stock)):
        bot_qty, ff_qty = balances.get(article), ff_stock.get(article)
        diff = (bot_qty - ff_qty) if (bot_qty is not None and ff_qty is not None) else "—"
        ws2.append([article, bot_qty if bot_qty is not None else "—", ff_qty if ff_qty is not None else "—", diff])
        if isinstance(diff, int) and diff != 0:
            ws2.cell(ws2.max_row, 4).fill = RED
    ws2.freeze_panes = "A2"
    for col, width in zip("ABCD", (18, 16, 22, 14)):
        ws2.column_dimensions[col].width = width

    # --- Расхождения (с комментарием, кто именно недосчитал) ---
    ws3 = wb.create_sheet("Расхождения")
    ws3.append(["Артикул", "Дата", "Площадка", "Схема", "Факт (API)", "Журнал ФФ", "Разница", "Комментарий"])
    for c in ws3[1]:
        c.font = BOLD
    for article, date, platform, scheme, api_qty, j_qty in mismatches:
        row = [article, date.isoformat(), PLATFORM_LABELS.get(platform, platform), scheme, api_qty, j_qty,
               api_qty - j_qty, _mismatch_comment(platform, api_qty, j_qty)]
        ws3.append(row)
        ws3.cell(ws3.max_row, 7).fill = RED
    ws3.freeze_panes = "A2"
    for col, width in zip("ABCDEFGH", (16, 12, 10, 8, 12, 12, 10, 70)):
        ws3.column_dimensions[col].width = width

    # --- Наш журнал отгрузок (то, что бот сам собрал по API, день за днём) ---
    ws4 = wb.create_sheet("Наш журнал отгрузок")
    ws4.append(["Дата", "Площадка", "Схема", "Артикул", "Название", "Количество"])
    for c in ws4[1]:
        c.font = BOLD
    for r in sorted(api_rows, key=lambda r: (r["date"], r["platform"], r["article"])):
        ws4.append([
            r["date"].isoformat(), PLATFORM_LABELS.get(r["platform"], r["platform"]), r["scheme"],
            r["article"], names_by_article.get(r["article"], ""), r["qty"],
        ])
    ws4.freeze_panes = "A2"
    for col, width in zip("ABCDEF", (12, 10, 8, 14, 50, 12)):
        ws4.column_dimensions[col].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
