import io

import openpyxl

from . import cabinets, margin, ozon_prices
from .ozon_client import OzonClient
from .wb_client import WBClient

HEADER = ["Артикул", "Название", "Себестоимость за шт (руб)"]


def _rows_for_cabinet(cabinet: dict):
    creds = cabinet["credentials"]
    existing = cabinets.get_cost_prices(cabinet["id"])

    if cabinet["marketplace"] == "ozon":
        client = OzonClient(creds["client_id"], creds["api_key"])
        items = ozon_prices.get_price_list(client=client)
        return [(i["offer_id"], i["name"], existing.get(i["offer_id"], "")) for i in items]

    # WB has no full-catalog endpoint wired up yet — fall back to whatever
    # sold in the last 60 days (from the margin build), which is a partial
    # list but covers everything actively contributing to profit.
    client = WBClient(creds["api_key"])
    summary = margin.build_margin_summary(client=client, cost_prices=existing, days=60)
    return [(str(p["nm_id"]), p["title"], existing.get(str(p["nm_id"]), "")) for p in summary["products"]]


def build_template(cabinet: dict) -> bytes:
    """Builds an .xlsx: Артикул | Название | Себестоимость (pre-filled with
    whatever cost prices are already stored). The user fills/edits column C
    and sends the file back to the bot."""
    rows = _rows_for_cabinet(cabinet)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Себестоимость"
    ws.append(HEADER)
    for key, name, cost in rows:
        ws.append([key, name, cost])
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 55
    ws.column_dimensions["C"].width = 26

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def parse_template(file_bytes: bytes) -> dict:
    """Reads the filled-in template back. Returns {item_key: cost_price} for
    every row with a positive number in column C — extra/reordered/manually
    added rows are fine as long as column A holds a valid item key."""
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    ws = wb.active
    prices = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        key = str(row[0]).strip()
        cost = row[2] if len(row) > 2 else None
        if key and isinstance(cost, (int, float)) and cost > 0:
            prices[key] = float(cost)
    return prices
